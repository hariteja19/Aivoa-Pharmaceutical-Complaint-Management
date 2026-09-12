import io
import pypdf
import pdfplumber
import docx
from app.core.exceptions import DocumentParsingException
from app.core.logging import logger

class DocumentParserService:
    @staticmethod
    def parse_pdf(file_bytes: bytes) -> str:
        extracted_text = ""
        try:
            # Try pdfplumber first for better structure/tables
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        extracted_text += text + "\n"
        except Exception as e:
            logger.warning(f"pdfplumber failed: {e}. Falling back to PyPDF.")
            try:
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        extracted_text += text + "\n"
            except Exception as ex:
                logger.error(f"PyPDF extraction error: {ex}")
                raise DocumentParsingException(f"Failed to extract text from PDF: {str(ex)}")

        if not extracted_text.strip():
            raise DocumentParsingException("PDF contains no extractable text or is a scanned image.")
        return extracted_text.strip()

    @staticmethod
    def parse_docx(file_bytes: bytes) -> str:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        full_text.append(" | ".join(row_data))
            return "\n".join(full_text)
        except Exception as e:
            logger.error(f"DOCX extraction error: {e}")
            raise DocumentParsingException(f"Failed to extract text from DOCX document: {str(e)}")

    @staticmethod
    def parse_txt(file_bytes: bytes) -> str:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return file_bytes.decode("latin-1")
            except Exception as e:
                raise DocumentParsingException(f"Failed to decode text file: {str(e)}")

    @classmethod
    def extract_text(cls, filename: str, file_bytes: bytes) -> str:
        fn = filename.lower()
        if fn.endswith(".pdf"):
            return cls.parse_pdf(file_bytes)
        elif fn.endswith(".docx") or fn.endswith(".doc"):
            return cls.parse_docx(file_bytes)
        elif fn.endswith(".txt"):
            return cls.parse_txt(file_bytes)
        else:
            raise DocumentParsingException(f"Unsupported file format: '{filename}'. Allowed: PDF, DOCX, TXT")
